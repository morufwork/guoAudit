"""Create species-stratified tables/figure and insert new sections without rebuilding existing DOCX content."""
from pathlib import Path
from shutil import copy2
from io import BytesIO
import importlib.util
import json
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/species_stratified_replication'
NAMES={'celegans':'C. elegans','dmelanogaster':'D. melanogaster','ecoli':'E. coli'}
VAR={'n0':'N0','n2':'N2-decile','n2_exact':'N2-exact','n2_nn':'N2-nearest'}
REPS={'degree_only':'Degree','aac_ctd':'AAC+CTD','esm2_mean':'ESM-2 mean'}
MOD={'logistic_regression':'LR','mlp':'MLP'}

def qualify_pooled_claims():
    replacements={
        'A separately constructed multi-species benchmark also exhibited degree-associated bias, yet ESM-2–MLP retained ROC-AUC 0.698 after degree matching.':
        'A separately constructed multi-species benchmark also exhibited degree-associated bias. Its pooled N2 result was additionally confounded by cross-species negative pairing; within-species replication is reported separately.',
        '3.7 A second benchmark retains sequence signal after degree control':
        '3.7 Pooled second-benchmark discrimination after degree control',
        'Sequence models, however, retained meaningful N2 performance:':
        'Sequence models retained pooled N2 discrimination:',
        'Thus, degree-associated bias can coexist with genuine residual sequence signal.':
        'However, the taxonomic audit found that 45.1% of pooled N2 negatives joined different species, and a same-species indicator alone achieved ROC-AUC 0.729. The pooled sequence result therefore cannot establish molecular interaction signal independent of taxonomic pairing cues. Section 3.8 reports the within-species controls.',
        'Crucially, sequence models retained ROC-AUC values well above chance under matched negatives. This result prevents a universal conclusion from a single benchmark failure.':
        'Sequence models retained pooled discrimination under matched negatives, but cross-species negative pairing provided an additional cue. The within-species analyses in Section 3.8 refine this comparison; their single-split estimates do not establish unseen-protein generalization or molecular compatibility.',
        'Protein language models, handcrafted sequence descriptors, and deep architectures may retain genuine utility when evaluated on datasets whose negative construction does not erase the relevant signal.':
        'Neither the Guo collapse nor pooled residual discrimination supports a universal claim about protein language models, handcrafted sequence descriptors, or deep architectures.',
    }
    changes={}
    path=ROOT/'manuscript/final_manuscrip_v2.docx';doc=Document(path)
    for p in doc.paragraphs:
        old=p.text;new=old
        for source,target in replacements.items():new=new.replace(source,target)
        if old!=new:changes[old]=new;p.text=new
    doc.save(path)
    md=ROOT/'manuscript/FINAL_MANUSCRIPT_V2.md';text=md.read_text()
    for source,target in replacements.items():
        assert source in text,source
        text=text.replace(source,target)
    md.write_text(text)
    (OUT/'manuscript_claim_revisions.json').write_text(json.dumps(changes,indent=2))

def table(doc,df):
    t=doc.add_table(rows=1,cols=len(df.columns));t.style='Table Grid';t.autofit=False
    width=6.6/len(df.columns)
    for col in t.columns:col.width=Inches(width)
    for c,v in zip(t.rows[0].cells,df.columns):c.text=str(v)
    header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
    for row in df.itertuples(index=False,name=None):
        for c,v in zip(t.add_row().cells,row):c.text=f'{v:.4f}' if isinstance(v,float) else str(v)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            c.width=Inches(width)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs:r.font.size=Pt(8);r.bold=i==0
    return t

def render_species_figure(metrics):
    def auc(s,v,r='esm2_mean',m='mlp'):
        d=metrics[(metrics.species==s)&(metrics.variant==v)&(metrics.representation==r)&(metrics.model==m)]
        assert len(d)==1;return d.iloc[0].roc_auc
    fig,axes=plt.subplots(1,3,figsize=(10,4.1),sharey=True)
    combinations=[('degree_only','logistic_regression'),('aac_ctd','logistic_regression'),('aac_ctd','mlp'),('esm2_mean','logistic_regression'),('esm2_mean','mlp')]
    labels=['Degree\nLR','AAC+\nCTD\nLR','AAC+\nCTD\nMLP','ESM-2\nLR','ESM-2\nMLP']
    for ax,s in zip(axes,NAMES):
        for v,color,offset in [('n0','#4c78a8',-.18),('n2','#f58518',.18)]:
            ax.bar([i+offset for i in range(5)],[auc(s,v,r,m) for r,m in combinations],width=.35,color=color,label=VAR[v])
        ax.set_xticks(range(5),labels,fontsize=9);ax.set_title(NAMES[s]);ax.set_ylim(0,1);ax.axhline(.5,color='gray',ls='--',lw=1)
        ax.spines[['top','right']].set_visible(False)
    axes[0].set_ylabel('ROC-AUC')
    handles,legend_labels=axes[-1].get_legend_handles_labels()
    fig.legend(handles,legend_labels,loc='upper center',ncol=2,frameon=False,fontsize=9)
    fig.tight_layout(rect=(0,0,1,.90));fig.savefig(ROOT/'figures/figure11_species_stratified.png',dpi=220);fig.savefig(ROOT/'figures/figure11_species_stratified.pdf');plt.close(fig)

def main():
    spec=importlib.util.spec_from_file_location('species_validation',ROOT/'scripts/52_validate_species_replication.py')
    validator=importlib.util.module_from_spec(spec);spec.loader.exec_module(validator)
    validator.main(data_only=True)
    dfs={k:pd.concat([pd.read_csv(OUT/s/f'{k}.csv') for s in NAMES],ignore_index=True) for k in ['metrics','balance','splits','degree_distributions']}
    for k,df in dfs.items():df.to_csv(ROOT/f'tables/species_stratified_{k}.csv',index=False)
    metrics=dfs['metrics'];balance=dfs['balance'];splits=dfs['splits']
    assert len(metrics)==60 and len(balance)==12 and len(splits)==12
    audit=pd.DataFrame([json.loads((OUT/s/'audit.json').read_text()) for s in NAMES]);audit.to_csv(ROOT/'tables/species_stratified_audit.csv',index=False)
    render_species_figure(metrics)
    accounting=pd.read_csv(ROOT/'tables/species_pair_accounting.csv').set_index('category')
    omitted=accounting.loc[~accounting.index.isin(NAMES)].sum()
    mapping=pd.read_csv(ROOT/'data/external/multi_species_taxonomy/species_assignments.tsv',sep='\t')
    pooled=pd.read_csv(ROOT/'tables/pooled_n2_species_accounting.csv').set_index('category').n_pairs
    n_cross=int(pooled.get('cross_species',0));n_pooled=int(pooled.sum())
    taxdiag=pd.read_csv(ROOT/'tables/pooled_same_species_diagnostic.csv')
    tax_n2=taxdiag.set_index('variant').loc['n2']
    paragraphs=[]
    sizes='; '.join(f"{NAMES[r.species]}, {r.n_proteins:,} proteins and {r.n_pairs:,} pairs ({r.n_positive:,} positive and {r.n_negative:,} N0)" for r in audit.itertuples())
    paragraphs.append('To determine whether the pooled result persisted within individual species, we separated the benchmark using accession-level taxonomy and repeated the combined benchmark protocol independently in C. elegans, D. melanogaster, and E. coli. The resulting datasets comprised '+sizes+'. Taxonomy resolution and exclusions are detailed in Supplementary Section S20 and Tables S25–S26. Positive-network degrees and matching strata were recomputed within each species, and all replacement endpoints were restricted to that species. Each analysis used the same 80/20 random pair split procedure, seed 42, frozen representations, symmetric pair fusion, and model settings as the combined analysis.')
    paragraphs.append(f"Reconstruction of the combined benchmark’s N2 sampler identified {n_cross:,} cross-species negative pairs among {n_pooled:,} generated negatives ({100*n_cross/n_pooled:.1f}%; Supplementary Table S32). The species-stratified N2 datasets contain only within-species pairs. A same-species indicator, using no sequence or degree features, achieved ROC-AUC {tax_n2.roc_auc:.3f} on {int(tax_n2.n_evaluated):,} pooled N2 test pairs whose endpoints were assigned to the three target species (Table S33). Thus, taxonomic pairing cues can contribute to the pooled N2 discrimination and must be considered alongside degree-associated effects.")
    for s in NAMES:
        paragraphs.append(f"In {NAMES[s]}, degree-only logistic regression changed from ROC-AUC {auc(s,'n0','degree_only','logistic_regression'):.3f} under N0 to {auc(s,'n2','degree_only','logistic_regression'):.3f} under N2. AAC+CTD–MLP changed from {auc(s,'n0','aac_ctd'):.3f} to {auc(s,'n2','aac_ctd'):.3f}, and ESM-2–MLP from {auc(s,'n0'):.3f} to {auc(s,'n2'):.3f}. With exact-degree and nearest-degree matching, ESM-2–MLP reached {auc(s,'n2_exact'):.3f} and {auc(s,'n2_nn'):.3f}, respectively; the corresponding degree-only values were {auc(s,'n2_exact','degree_only','logistic_regression'):.3f} and {auc(s,'n2_nn','degree_only','logistic_regression'):.3f}.")
    exact=balance[balance.variant=='n2_exact'];nearest=balance[balance.variant=='n2_nn']
    paragraphs.append('Exact-degree substitution skipped '+', '.join(f'{int(r.skipped):,} positives in {NAMES[r.species]}' for r in exact.itertuples())+'. Exact matching therefore does not guarantee balance against the complete positive set: unmatchable positives remain in the evaluation while their proposed negatives are omitted. Nearest-degree matching retained one negative per positive; its endpoint-degree SMDs were '+', '.join(f'{r.standardized_mean_difference:.3f} in {NAMES[r.species]}' for r in nearest.itertuples())+'. These realized balance and class-composition differences must accompany interpretation of residual discrimination.')
    # Interpretation is written after checking the generated values; avoid an automatic above-chance claim.
    paragraphs.append('Figure 11 and Supplementary Tables S27–S30 report the species-specific model and degree-balance results. These are descriptive estimates from a single random pair split per species, with proteins permitted to recur across training and test pairs. They do not establish generalization to unseen proteins or statistically significant differences between species. Separating species also changes candidate pools, degree strata, and negative-set composition, so the pooled estimate is not an average of these independently reconstructed evaluations.')
    heading='3.8 Species-stratified replication of degree-control analyses'
    caption='Figure 11. Species-stratified replication of the multi-species benchmark. ROC-AUC under original (N0) and within-species degree-decile-matched (N2) negatives for degree-only logistic regression and AAC+CTD/ESM-2 mean representations with logistic regression (LR) or MLP. Each species uses independently reconstructed random pair splits (80/20; seed 42). Bars are point estimates, without inferential error bars. Exact-degree and nearest-degree sensitivity results and all plotted values appear in Supplementary Tables S27–S29.'
    mainpath=ROOT/'manuscript/final_manuscrip_v2.docx';doc=Document(mainpath)
    assert not any(p.text==heading for p in doc.paragraphs)
    copy2(mainpath,OUT/'final_manuscrip_v2_before_species.docx')
    anchor=next(p for p in doc.paragraphs if p.text=='4 Discussion')
    def insert(p):anchor._p.addprevious(p._p)
    insert(doc.add_heading(heading,level=3))
    for text in paragraphs:insert(doc.add_paragraph(text))
    p=doc.add_paragraph();p.add_run().add_picture(str(ROOT/'figures/figure11_species_stratified.png'),width=Inches(6.6));insert(p)
    p=doc.add_paragraph(caption);insert(p)
    doc.save(mainpath)
    mdpath=ROOT/'manuscript/FINAL_MANUSCRIPT_V2.md'
    section='### '+heading+'\n\n'+'\n\n'.join(paragraphs)+'\n\n[FIGURE:figures/figure11_species_stratified.pdf]\n\n**'+caption+'**\n\n'
    mdpath.write_text(mdpath.read_text().replace('## 4 Discussion',section+'## 4 Discussion',1))
    spath=ROOT/'manuscript/supplementary_v2.docx';sup=Document(spath);copy2(spath,OUT/'supplementary_v2_before_species.docx')
    sup.add_page_break();sh='S20 Species-stratified benchmark audit and replication';sup.add_heading(sh,level=1)
    method='Taxonomy was resolved from cached UniProtKB accession records and NCBI protein ESummary records for legacy GI identifiers. Retired UniProt accessions were checked against UniSave version 1; archived organism names were used when those early records lacked a numerical taxon identifier. Unresolved labels were assigned by exact sequence identity only when all matching resolved benchmark records belonged to one species. E. coli strain records were consolidated. Proteins belonging to other organisms, unresolved endpoints, and cross-species pairs were excluded from the three species-specific datasets; the original combined results were retained.'
    protocol='Within each species, the existing canonicalized, self-pair-filtered input was partitioned without changing its sequences or original labels. Degrees were computed from that species’s full positive graph, including zero-degree proteins in the candidate universe. Negative sets used the original N0 labels or the same one-sided decile, exact-degree, and nearest-degree substitution functions as scripts 42–43. Self-pairs, known positives, and duplicate negatives were rejected. Unmatchable substitutions were skipped, with realized class counts reported. Positive pairs were split once (80/20; random_state=42); each negative set was independently split with the same settings. The positive split was shared across variants and all five models used identical realized pairs within each variant. N2 negatives were used consistently for both training and test.'
    training='AAC+CTD features and frozen ESM-2 mean embeddings were reused from the combined analysis. Pair fusion concatenated the sum, absolute difference, and elementwise product. StandardScaler was fitted on training data only. Sequence logistic regression used max_iter=2000; the MLP used hidden layers (128,64), max_iter=500, and early_stopping=True, with the existing sklearn validation defaults. Degree-only logistic regression used the same default settings as script 42. All estimators used random_state=42. ROC-AUC, average precision (reported as PR-AUC), MCC, F1, balanced accuracy, accuracy, precision, recall, and specificity were saved; threshold-based metrics used 0.5. No per-species hyperparameter tuning was performed.'
    caution='Degree SMD, Kolmogorov–Smirnov (KS) distance, and Wasserstein distance describe pooled endpoint occurrences; they are not independent-endpoint significance tests. The existing 2,000-replicate distinct-protein bootstrap diagnostic was also retained in the machine-readable balance output for protocol parity, but is not used here to claim model significance. No cluster-aware model comparisons or multi-seed analyses were added to this single-split replication. Full-graph degree is a retrospective benchmark diagnostic. Random pair splits permit protein and sequence recurrence, so residual discrimination cannot be attributed exclusively to molecular compatibility.'
    exclusions=f"Of {len(mapping):,} input proteins, {(mapping.species=='other').sum():,} were assigned to other organisms and {(mapping.species=='unresolved').sum():,} remained unresolved. The excluded pair categories comprised {int(omitted['positive']):,} positive and {int(omitted['negative']):,} N0 pairs in total (Table S26); all retained and excluded counts reconcile with the combined cleaned input. Accession-level assignments and sources are recorded in data/external/multi_species_taxonomy/species_assignments.tsv."
    for p in [method,exclusions,protocol,training,caution]:sup.add_paragraph(p)
    smd=['\n## '+sh+'\n',method,exclusions,protocol,training,caution]
    tables=[]
    a=audit[['species','n_proteins','n_positive','n_negative','n_pairs','zero_positive_degree']].copy();a.species=a.species.map(NAMES);a.columns=['Species','Proteins','Positive pairs','N0 pairs','Total pairs','Zero-degree proteins'];tables.append((25,'Dataset dimensions after taxonomy-based separation.',a))
    acc=accounting.reset_index();acc.columns=['Pair category','Negative pairs','Positive pairs'];tables.append((26,'Accounting for all pairs in the combined cleaned benchmark.',acc))
    for num,s in zip([27,28,29],NAMES):
        d=metrics[metrics.species==s].copy();d['Set']=d.variant.map(VAR);d['Representation']=d.representation.map(REPS);d['Model']=d.model.map(MOD)
        d=d[['Set','Representation','Model','n_train','n_test','roc_auc','pr_auc','mcc']];d.columns=['Set','Representation','Model','Train pairs','Test pairs','ROC-AUC','PR-AUC','MCC'];tables.append((num,f'{NAMES[s]} model comparison and matching-granularity sensitivity.',d))
    b=balance[['species','variant','n_negatives','skipped','standardized_mean_difference','ks_statistic','wasserstein_distance']].copy();b.species=b.species.map(NAMES);b.variant=b.variant.map(VAR);b.columns=['Species','Set','Negatives','Skipped','SMD','KS distance','Wasserstein'];tables.append((30,'Within-species endpoint-degree balance across negative sets.',b))
    d=splits[['species','variant','train_positive','train_negative','test_positive','test_negative','overlapping_proteins']].copy();d.species=d.species.map(NAMES);d.variant=d.variant.map(VAR);d.columns=['Species','Set','Train +','Train −','Test +','Test −','Shared proteins'];tables.append((31,'Realized split counts and train/test protein recurrence.',d))
    pooled_table=pooled.rename('Negative pairs').reset_index().rename(columns={'category':'Pair category'})
    pooled_table['Percent']=100*pooled_table['Negative pairs']/n_pooled
    tables.append((32,'Taxonomic pair composition of the reconstructed pooled N2 negative set.',pooled_table))
    td=taxdiag[['variant','n_evaluated','n_excluded','roc_auc','pr_auc','mcc']].copy();td.columns=['Set','Test pairs','Excluded pairs','ROC-AUC','PR-AUC','MCC']
    tables.append((33,'Same-species indicator diagnostic on pooled random-split test pairs.',td))
    for num,title,df in tables:
        csv=f'tables/tableS{num}_species_replication.csv';df.to_csv(ROOT/csv,index=False)
        sup.add_page_break();p=sup.add_paragraph();p.add_run(f'Supplementary Table S{num}. {title}').bold=True;p.paragraph_format.keep_with_next=True
        table(sup,df);note='Values are rounded to four decimal places where applicable. LR: logistic regression. N2-decile, N2-exact, and N2-nearest use degree-decile, exact-degree, and nearest-degree matching, respectively. All pair counts refer to the realized datasets.'
        sup.add_paragraph(note);sup.add_paragraph('Source: '+csv).runs[0].italic=True
        smd.extend([f'**Supplementary Table S{num}. {title}**',f'[CSV:{csv}]',note])
    provenance='Reproduction: scripts/48_fetch_multi_species_taxonomy.py caches taxonomy; scripts/50_audit_species_assignments.py assigns species and audits pair categories; scripts/49_species_stratified_replication.py --species celegans (or dmelanogaster/ecoli) runs each analysis; scripts/51_report_species_replication.py generates the new reporting sections. Full metrics, degree summaries, and balance diagnostics are in tables/species_stratified_*.csv. Per-species inputs, splits, predictions, fit warnings, and environment manifests are in results/species_stratified_replication/. Taxonomy sources: https://rest.uniprot.org/ and https://eutils.ncbi.nlm.nih.gov/entrez/eutils/.'
    taxnote='The same-species diagnostic assigns score 1 when both endpoints belong to the same target species and score 0 otherwise. It uses no fitted classifier. Evaluation uses the reconstructed combined N0/N2 test partitions (seed 42), retaining only pairs with both endpoints assigned to C. elegans, D. melanogaster, or E. coli; excluded counts are reported in Table S33. This is a retrospective construction diagnostic, not a deployable PPI predictor.'
    sup.add_paragraph(taxnote);smd.append(taxnote)
    sup.add_paragraph(provenance);smd.append(provenance);sup.save(spath)
    with (ROOT/'manuscript/SUPPLEMENTARY_INFORMATION_V2.md').open('a') as f:f.write('\n\n'.join(smd)+'\n')
    qualify_pooled_claims()
    print('Created manuscript section 3.8, Figure 11, supplementary section S20, and Tables S25–S33.')
if __name__=='__main__':main()
